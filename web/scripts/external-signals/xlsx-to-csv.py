#!/usr/bin/env python3
"""xlsx-to-csv.py <in.xlsx> <out.csv> [sheet-index]

Converts one worksheet of an .xlsx to UTF-8 CSV (G14-S40). Used by
`ingest.mjs --fetch --source rdti-transparency` for the ATO R&D Tax
Incentive transparency report, whose data sheet is the second sheet (the
first is a cover note). Default sheet = the one with the most rows.
Requires openpyxl (present on the production host: `python3 -c "import openpyxl"`).
"""
import csv
import sys

import openpyxl


def main() -> int:
    if len(sys.argv) < 3:
        print(__doc__)
        return 2
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    if len(sys.argv) > 3:
        ws = wb.worksheets[int(sys.argv[3])]
    else:
        ws = max(wb.worksheets, key=lambda s: s.max_row or 0)
    n = 0
    with open(dst, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        for row in ws.iter_rows(values_only=True):
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            w.writerow(["" if v is None else v for v in row])
            n += 1
    print(f"{dst}: {n} rows from sheet '{ws.title}'")
    return 0


if __name__ == "__main__":
    sys.exit(main())
